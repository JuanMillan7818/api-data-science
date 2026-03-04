#!/usr/bin/env python3
import pandas as pd
import numpy as np
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.comments import Comment

def generar_outliers_excel():
    """Genera automáticamente Excel con outliers marcados para centenarios_completo_preprocesado.xlsx"""
    
    print("🔍 Generando Excel con outliers marcados...")
    
    # Cargar datos
    dataset_path = Path("centenarios/centenarios_completo_preprocesado.xlsx")
    
    if not dataset_path.exists():
        print(f"❌ Error: No existe el archivo {dataset_path}")
        return
    
    try:
        # Cargar con header=1
        df = pd.read_excel(dataset_path, header=1)
        print(f"📊 Datos cargados: {df.shape[0]} filas × {df.shape[1]} columnas")
        
        # Identificar columnas numéricas
        numeric_cols = df.select_dtypes(include=[np.number]).columns
        print(f"🔢 Columnas numéricas: {len(numeric_cols)}")
        
        # Excluir columnas de ID
        exclude_cols = ["CEDULA", "Número", "Cedula", "Numero"]
        numeric_cols = [col for col in numeric_cols if col not in exclude_cols]
        
        # Detectar outliers
        outliers_info = {}
        
        for col in numeric_cols:
            Q1 = df[col].quantile(0.25)
            Q3 = df[col].quantile(0.75)
            IQR = Q3 - Q1
            
            lower_bound = Q1 - 1.5 * IQR
            upper_bound = Q3 + 1.5 * IQR
            
            outliers = df[(df[col] < lower_bound) | (df[col] > upper_bound)]
            
            if len(outliers) > 0:
                outliers_info[col] = {
                    'outliers_df': outliers,
                    'lower_bound': lower_bound,
                    'upper_bound': upper_bound
                }
        
        print(f"⚠️  Variables con outliers: {len(outliers_info)}")
        
        if len(outliers_info) == 0:
            print("✅ No se detectaron outliers")
            return
        
        # Guardar Excel base
        output_file = "centenarios_completo_preprocesado_con_outliers.xlsx"
        df.to_excel(output_file, index=False)
        
        # Cargar con openpyxl para marcar outliers
        wb = load_workbook(output_file)
        ws = wb.active
        
        # Definir estilos
        red_fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid")
        red_font = Font(color="FFFFFF", bold=True)
        yellow_font = Font(color="000000", bold=True)
        
        total_marked = 0
        
        # Marcar outliers
        for col_name, info in outliers_info.items():
            # Encontrar columna en Excel
            col_idx = None
            for cell in ws[1]:  # Primera fila (headers)
                if cell.value == col_name:
                    col_idx = cell.column
                    break
            
            if col_idx is None:
                print(f"⚠️ Columna '{col_name}' no encontrada en el Excel")
                continue
            
            # Marcar cada outlier
            for idx, row in info['outliers_df'].iterrows():
                excel_row = idx + 2  # +2 por header y 0-indexing
                
                cell_value = ws.cell(row=excel_row, column=col_idx).value
                
                if cell_value is not None:
                    # Determinar tipo y color
                    if cell_value < info['lower_bound']:
                        ws.cell(row=excel_row, column=col_idx).fill = yellow_fill
                        ws.cell(row=excel_row, column=col_idx).font = yellow_font
                        outlier_type = "BAJO"
                    elif cell_value > info['upper_bound']:
                        ws.cell(row=excel_row, column=col_idx).fill = red_fill
                        ws.cell(row=excel_row, column=col_idx).font = red_font
                        outlier_type = "ALTO"
                    
                    # Agregar comentario
                    comment_text = f"OUTLIER {outlier_type}\n"
                    comment_text += f"Valor: {cell_value}\n"
                    comment_text += f"Rango normal: [{info['lower_bound']:.2f}, {info['upper_bound']:.2f}]"
                    
                    cell = ws.cell(row=excel_row, column=col_idx)
                    if not cell.comment:
                        cell.comment = Comment(comment_text, "OutlierDetector")
                    
                    total_marked += 1
        
        # Guardar archivo modificado
        wb.save(output_file)
        
        print(f"✅ Excel con outliers guardado en: {output_file}")
        print(f"🔴 Rojo: Outliers por encima del rango normal")
        print(f"🟡 Amarillo: Outliers por debajo del rango normal")
        print(f"💬 Comentarios: Pasa el mouse sobre las celdas marcadas")
        print(f"📊 Total celdas marcadas: {total_marked}")
        
        # Resumen por variable
        print(f"\n📋 Resumen de outliers por variable:")
        for col_name, info in outliers_info.items():
            count = len(info['outliers_df'])
            print(f"  • {col_name}: {count} outliers")
        
    except Exception as e:
        print(f"❌ Error: {str(e)}")

if __name__ == "__main__":
    generar_outliers_excel()
